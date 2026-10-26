import io
import json
import zipfile

import docx
import pytest
from openpyxl import Workbook, load_workbook
from reportlab.lib.pagesizes import A4
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table
from reportlab.lib.styles import getSampleStyleSheet

import converters
from app import app
from converters import ConversionError

ROWS = [["Name", "Qty", "Zip"], ["Apple", 3, "02134"], ["Pear", 4.5, "10001"]]


def make_xlsx(sheets=None):
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in (sheets or {"Fruit": ROWS}).items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_csv():
    return b"Name,Qty,Zip\nApple,3,02134\nPear,4.5,10001\n"


def make_docx():
    d = docx.Document()
    d.add_heading("Report", 1)
    d.add_paragraph("Hello world")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "A", "B"
    t.cell(1, 0).text, t.cell(1, 1).text = "1", "2"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_pdf():
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4)
    doc.build([
        Paragraph("Quarterly summary text", getSampleStyleSheet()["BodyText"]),
        Table([["Name", "Qty"], ["Apple", "3"]], style=[("GRID", (0, 0), (-1, -1), 1, "black")]),
    ])
    return buf.getvalue()


SAMPLES = {"xlsx": make_xlsx, "csv": make_csv, "docx": make_docx, "pdf": make_pdf}


def fake_fetch(url):
    kind, _ = converters.parse_google_url(url)
    return ("docx", make_docx()) if kind == "gdoc" else ("xlsx", make_xlsx())


def _cases():
    for src in SAMPLES:
        for dst in converters.outputs_for(src):
            yield src, dst
    for src in ("gdoc", "gsheet"):
        for dst in converters.outputs_for(src):
            yield src, dst


HAS_LO = converters.libreoffice_path() is not None


@pytest.mark.parametrize("src,dst", list(_cases()))
def test_every_conversion(src, dst):
    if dst == "doc" and not HAS_LO:
        pytest.skip("LibreOffice not installed")
    if src == "gdoc":
        data = b"https://docs.google.com/document/d/abc123/edit"
    elif src == "gsheet":
        data = json.dumps({"url": "https://docs.google.com/spreadsheets/d/xyz/edit"}).encode()
    else:
        data = SAMPLES[src]()
    res = converters.convert(data, src, dst, stem="out", fetch=fake_fetch)
    assert res.data
    assert res.filename.startswith("out.")
    if dst == "pdf":
        assert res.data.startswith(b"%PDF")
    if dst in ("xlsx", "docx"):
        assert zipfile.is_zipfile(io.BytesIO(res.data))
    if dst == "json":
        json.loads(res.data)


def test_csv_to_xlsx_keeps_numbers_and_leading_zeros():
    res = converters.convert(make_csv(), "csv", "xlsx")
    ws = load_workbook(io.BytesIO(res.data)).active
    assert ws["B2"].value == 3
    assert ws["B3"].value == 4.5
    assert ws["C2"].value == "02134"


def test_xlsx_to_csv_roundtrip():
    res = converters.convert(make_xlsx(), "xlsx", "csv")
    text = res.data.decode("utf-8-sig")
    assert text.splitlines() == ["Name,Qty,Zip", "Apple,3,02134", "Pear,4.5,10001"]


def test_multi_sheet_csv_is_zip():
    res = converters.convert(make_xlsx({"A": ROWS, "B": ROWS}), "xlsx", "csv", stem="book")
    assert res.filename == "book.zip"
    names = zipfile.ZipFile(io.BytesIO(res.data)).namelist()
    assert names == ["book - A.csv", "book - B.csv"]


def test_xlsx_to_json_records():
    res = converters.convert(make_xlsx(), "xlsx", "json")
    assert json.loads(res.data)[0] == {"Name": "Apple", "Qty": 3, "Zip": "02134"}


def test_pdf_table_to_xlsx():
    res = converters.convert(make_pdf(), "pdf", "xlsx")
    ws = load_workbook(io.BytesIO(res.data)).active
    assert [c.value for c in ws[1]] == ["Name", "Qty"]
    assert [c.value for c in ws[2]] == ["Apple", 3]


def test_docx_to_markdown():
    md = converters.convert(make_docx(), "docx", "md").data.decode()
    assert "# Report" in md and "Hello world" in md and "| A | B |" in md


def test_google_url_parsing():
    assert converters.parse_google_url("https://docs.google.com/document/u/0/d/AbC-1_x/edit") == ("gdoc", "AbC-1_x")
    assert converters.parse_google_url("https://docs.google.com/spreadsheets/d/Z9/edit#gid=0") == ("gsheet", "Z9")
    with pytest.raises(ConversionError):
        converters.parse_google_url("https://example.com/foo")


def test_unsupported_pair_and_type():
    with pytest.raises(ConversionError):
        converters.convert(make_csv(), "csv", "csv")
    with pytest.raises(ConversionError):
        converters.detect_format("photo.png")


# ---------------------------------------------------------------- HTTP layer


@pytest.fixture
def client():
    app.config["TESTING"] = True
    return app.test_client()


def test_index_and_formats(client):
    assert client.get("/").status_code == 200
    data = client.get("/api/formats").get_json()
    assert "xlsx" in data["inputs"]["csv"]["outputs"]


def test_api_convert_upload(client):
    resp = client.post(
        "/api/convert",
        data={"target": "xlsx", "file": (io.BytesIO(make_csv()), "sales data.csv")},
        content_type="multipart/form-data",
    )
    assert resp.status_code == 200
    assert "sales_data.xlsx" in resp.headers["Content-Disposition"]


def test_api_convert_errors(client):
    assert client.post("/api/convert", data={"target": "pdf"}).status_code == 400
    resp = client.post("/api/convert", data={"target": "pdf", "url": "https://example.com"})
    assert resp.status_code == 400 and "Google" in resp.get_json()["error"]

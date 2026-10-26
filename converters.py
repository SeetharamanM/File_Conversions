"""Conversion engine.

Every input is loaded into one of two intermediate shapes:

* a *workbook*: an ordered dict of ``sheet name -> list of rows`` (tabular data)
* a *document*: a list of blocks ``("heading", text, level)``, ``("para", text)``
  or ``("table", rows)``

Writers then render either shape into the requested output format. Where
LibreOffice is installed it is used for higher-fidelity Office -> PDF/DOC
conversions; otherwise the pure-Python writers are used.
"""

from __future__ import annotations

import csv
import html
import io
import json
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from collections import OrderedDict
from dataclasses import dataclass

import pandas as pd
import requests

# --------------------------------------------------------------------------
# Format catalogue
# --------------------------------------------------------------------------

SPREADSHEET_INPUTS = {"xlsx", "xls", "csv", "gsheet"}
DOCUMENT_INPUTS = {"docx", "doc", "gdoc"}
PDF_INPUTS = {"pdf"}

SPREADSHEET_OUTPUTS = ["xlsx", "csv", "pdf", "docx", "json", "html", "txt", "md"]
DOCUMENT_OUTPUTS = ["pdf", "docx", "doc", "txt", "html", "md", "xlsx", "csv"]
PDF_OUTPUTS = ["docx", "txt", "html", "md", "xlsx", "csv", "json"]

MIME_TYPES = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv",
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "doc": "application/msword",
    "json": "application/json",
    "html": "text/html",
    "txt": "text/plain",
    "md": "text/markdown",
    "zip": "application/zip",
}

LABELS = {
    "xlsx": "Excel (.xlsx)",
    "xls": "Excel 97-2003 (.xls)",
    "csv": "CSV",
    "pdf": "PDF",
    "docx": "Word (.docx)",
    "doc": "Word 97-2003 (.doc)",
    "gdoc": "Google Doc",
    "gsheet": "Google Sheet",
    "json": "JSON",
    "html": "HTML",
    "txt": "Plain text",
    "md": "Markdown",
}


class ConversionError(Exception):
    """Raised for user-facing conversion problems."""


def outputs_for(fmt: str) -> list[str]:
    if fmt in SPREADSHEET_INPUTS:
        outs = SPREADSHEET_OUTPUTS
    elif fmt in DOCUMENT_INPUTS:
        outs = DOCUMENT_OUTPUTS
    elif fmt in PDF_INPUTS:
        outs = PDF_OUTPUTS
    else:
        raise ConversionError(f"Unsupported input format: {fmt!r}")
    return [o for o in outs if o != fmt]


def format_catalogue() -> dict:
    inputs = sorted(SPREADSHEET_INPUTS | DOCUMENT_INPUTS | PDF_INPUTS)
    return {
        "inputs": {f: {"label": LABELS[f], "outputs": outputs_for(f)} for f in inputs},
        "labels": LABELS,
        "libreoffice": libreoffice_path() is not None,
    }


@dataclass
class Result:
    data: bytes
    filename: str
    mimetype: str


# --------------------------------------------------------------------------
# LibreOffice helper
# --------------------------------------------------------------------------


def libreoffice_path() -> str | None:
    return shutil.which("soffice") or shutil.which("libreoffice")


def libreoffice_convert(data: bytes, src_ext: str, dst_ext: str) -> bytes:
    soffice = libreoffice_path()
    if not soffice:
        raise ConversionError(
            f"Converting .{src_ext} to .{dst_ext} requires LibreOffice, which is not installed."
        )
    with tempfile.TemporaryDirectory() as tmp:
        src = os.path.join(tmp, f"input.{src_ext}")
        with open(src, "wb") as fh:
            fh.write(data)
        cmd = [
            soffice,
            f"-env:UserInstallation=file://{tmp}/profile",
            "--headless",
            "--convert-to",
            dst_ext,
            "--outdir",
            tmp,
            src,
        ]
        try:
            subprocess.run(cmd, capture_output=True, timeout=120, check=False)
        except subprocess.TimeoutExpired as exc:
            raise ConversionError("LibreOffice timed out while converting the file.") from exc
        out = os.path.join(tmp, f"input.{dst_ext}")
        if not os.path.exists(out):
            raise ConversionError(f"LibreOffice could not convert .{src_ext} to .{dst_ext}.")
        with open(out, "rb") as fh:
            return fh.read()


# --------------------------------------------------------------------------
# Google Docs / Sheets
# --------------------------------------------------------------------------

_GOOGLE_RE = re.compile(
    r"docs\.google\.com/(?P<kind>document|spreadsheets)/(?:u/\d+/)?d/(?P<id>[A-Za-z0-9_-]+)"
)


def parse_google_url(url: str) -> tuple[str, str]:
    """Return ``(fmt, file_id)`` where fmt is ``gdoc`` or ``gsheet``."""
    m = _GOOGLE_RE.search(url or "")
    if not m:
        raise ConversionError(
            "That doesn't look like a Google Docs or Google Sheets link "
            "(expected https://docs.google.com/document/d/... or /spreadsheets/d/...)."
        )
    return ("gdoc" if m.group("kind") == "document" else "gsheet"), m.group("id")


def google_url_from_shortcut(data: bytes) -> str:
    """Read the URL from a Google Drive ``.gdoc``/``.gsheet`` shortcut file."""
    try:
        info = json.loads(data.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ConversionError("Could not read the Google shortcut file.") from exc
    url = info.get("url")
    if not url and info.get("doc_id"):
        kind = "spreadsheets" if info.get("resource_id", "").startswith("spreadsheet") else "document"
        url = f"https://docs.google.com/{kind}/d/{info['doc_id']}"
    if not url:
        raise ConversionError("The Google shortcut file does not contain a document link.")
    return url


def fetch_google_file(url: str, session: requests.Session | None = None) -> tuple[str, bytes]:
    """Download a shared Google Doc as .docx or a Google Sheet as .xlsx.

    Returns ``(downloaded_ext, data)``. The file must be shared as
    "Anyone with the link can view".
    """
    fmt, file_id = parse_google_url(url)
    if fmt == "gdoc":
        ext, export = "docx", f"https://docs.google.com/document/d/{file_id}/export?format=docx"
    else:
        ext, export = "xlsx", f"https://docs.google.com/spreadsheets/d/{file_id}/export?format=xlsx"
    http = session or requests
    try:
        resp = http.get(export, timeout=60, allow_redirects=True)
    except requests.RequestException as exc:
        raise ConversionError(f"Could not reach Google: {exc}") from exc
    ctype = resp.headers.get("Content-Type", "")
    if resp.status_code != 200 or "text/html" in ctype:
        raise ConversionError(
            "Google refused the download. Make sure the file is shared as "
            "'Anyone with the link can view'."
        )
    return ext, resp.content


# --------------------------------------------------------------------------
# Readers
# --------------------------------------------------------------------------


def _clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and pd.isna(value):
        return ""
    if value is pd.NaT:
        return ""
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    return value


def _frame_rows(df: pd.DataFrame) -> list[list]:
    return [[_clean(v) for v in row] for row in df.itertuples(index=False, name=None)]


def read_workbook(data: bytes, fmt: str) -> "OrderedDict[str, list[list]]":
    try:
        if fmt == "csv":
            text = _decode_text(data)
            try:
                dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
                delim = dialect.delimiter
            except csv.Error:
                delim = ","
            rows = [row for row in csv.reader(io.StringIO(text), delimiter=delim)]
            return OrderedDict([("Sheet1", [[_coerce_number(c) for c in r] for r in rows])])
        engine = "xlrd" if fmt == "xls" else "openpyxl"
        frames = pd.read_excel(io.BytesIO(data), sheet_name=None, header=None, dtype=object, engine=engine)
    except ConversionError:
        raise
    except Exception as exc:  # pandas raises many different types
        raise ConversionError(f"Could not read the {fmt.upper()} file: {exc}") from exc
    return OrderedDict((name, _frame_rows(df)) for name, df in frames.items())


def _decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


_NUM_RE = re.compile(r"^-?\d+(\.\d+)?([eE][-+]?\d+)?$")


def _coerce_number(value: str):
    """Turn numeric-looking CSV strings into numbers, keeping leading zeros intact."""
    s = value.strip()
    if not _NUM_RE.match(s):
        return value
    digits = s.lstrip("-")
    if len(digits) > 1 and digits[0] == "0" and digits[1] != ".":
        return value  # e.g. zip codes, IDs
    if len(digits.replace(".", "")) > 15:
        return value  # beyond float precision; keep exact text
    try:
        return int(s) if re.fullmatch(r"-?\d+", s) else float(s)
    except ValueError:
        return value


def read_docx(data: bytes) -> list[tuple]:
    import docx
    from docx.document import Document as _Doc
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document: _Doc = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ConversionError(f"Could not read the Word file: {exc}") from exc

    blocks: list[tuple] = []
    body = document.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name if para.style is not None else "") or ""
            m = re.match(r"Heading (\d)", style)
            if style == "Title":
                blocks.append(("heading", text, 1))
            elif m:
                blocks.append(("heading", text, min(int(m.group(1)), 6)))
            elif style.startswith("List"):
                blocks.append(("para", "\u2022 " + text))
            else:
                blocks.append(("para", text))
        elif tag == "tbl":
            table = Table(child, document)
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if rows:
                blocks.append(("table", rows))
    return blocks


def read_pdf(data: bytes) -> list[tuple]:
    import pdfplumber

    blocks: list[tuple] = []
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for number, page in enumerate(pdf.pages, start=1):
                if len(pdf.pages) > 1:
                    blocks.append(("heading", f"Page {number}", 2))
                tables = page.find_tables()
                # Text outside the detected tables, so table content isn't duplicated.
                remaining = page
                for t in tables:
                    remaining = remaining.outside_bbox(t.bbox)
                text = remaining.extract_text() or ""
                for chunk in re.split(r"\n\s*\n", text):
                    chunk = " ".join(line.strip() for line in chunk.splitlines() if line.strip())
                    if chunk:
                        blocks.append(("para", chunk))
                for t in tables:
                    rows = [[(c or "").strip() for c in row] for row in t.extract()]
                    if rows:
                        blocks.append(("table", rows))
    except Exception as exc:
        raise ConversionError(f"Could not read the PDF: {exc}") from exc
    if not any(b[0] in ("para", "table") for b in blocks):
        raise ConversionError(
            "No text found in this PDF. It may be a scanned image, which needs OCR."
        )
    return blocks


def pdf_line_rows(data: bytes) -> "OrderedDict[str, list[list]]":
    """Fallback tabular view of a PDF with no tables: one text line per row."""
    import pdfplumber

    sheets: OrderedDict[str, list[list]] = OrderedDict()
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for number, page in enumerate(pdf.pages, start=1):
            lines = [l for l in (page.extract_text() or "").splitlines() if l.strip()]
            if lines:
                sheets[f"Page {number}"] = [[l] for l in lines]
    return sheets


# --------------------------------------------------------------------------
# Shape adapters
# --------------------------------------------------------------------------


def workbook_to_blocks(book: "OrderedDict[str, list[list]]") -> list[tuple]:
    blocks: list[tuple] = []
    multi = len(book) > 1
    for name, rows in book.items():
        if multi:
            blocks.append(("heading", name, 2))
        rows = _trim_rows(rows)
        if rows:
            blocks.append(("table", [[_cell_text(c) for c in r] for r in rows]))
        else:
            blocks.append(("para", "(empty sheet)"))
    return blocks


def blocks_to_workbook(blocks: list[tuple]) -> "OrderedDict[str, list[list]]":
    """Tables become sheets; if there are none, the text goes one paragraph per row."""
    tables = [b[1] for b in blocks if b[0] == "table"]
    if tables:
        return OrderedDict(
            (f"Table {i}", [[_coerce_number(c) for c in r] for r in rows])
            for i, rows in enumerate(tables, start=1)
        )
    rows = [[b[1]] for b in blocks if b[0] in ("heading", "para")]
    return OrderedDict([("Text", rows)])


def _trim_rows(rows: list[list]) -> list[list]:
    """Drop fully-empty trailing rows and columns."""
    rows = [list(r) for r in rows]
    while rows and all(_cell_text(c) == "" for c in rows[-1]):
        rows.pop()
    width = max((len(r) for r in rows), default=0)
    while width and all(len(r) < width or _cell_text(r[width - 1]) == "" for r in rows):
        width -= 1
    return [(r + [""] * width)[:width] for r in rows]


def _cell_text(value) -> str:
    value = _clean(value)
    if hasattr(value, "isoformat") and not isinstance(value, str):
        text = value.isoformat(sep=" ") if hasattr(value, "hour") else value.isoformat()
        return text[:-9] if text.endswith(" 00:00:00") else text
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


# --------------------------------------------------------------------------
# Writers: workbook
# --------------------------------------------------------------------------


def _safe_sheet_name(name: str, used: set) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "_", str(name))[:31] or "Sheet"
    candidate, n = base, 2
    while candidate.lower() in used:
        suffix = f" ({n})"
        candidate = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(candidate.lower())
    return candidate


def write_xlsx(book: "OrderedDict[str, list[list]]") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()
    for name, rows in book.items():
        ws = wb.create_sheet(_safe_sheet_name(name, used))
        for row in rows:
            ws.append([_clean(v) for v in row])
        widths: dict[int, int] = {}
        for row in rows:
            for i, v in enumerate(row, start=1):
                widths[i] = max(widths.get(i, 0), min(len(_cell_text(v)), 60))
        for i, w in widths.items():
            ws.column_dimensions[get_column_letter(i)].width = max(w + 2, 8)
        if rows:
            for cell in ws[1]:
                cell.font = Font(bold=True)
    if not wb.sheetnames:
        wb.create_sheet("Sheet1")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _csv_bytes(rows: list[list]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\n")
    for r in rows:
        writer.writerow([_cell_text(c) for c in r])
    return buf.getvalue().encode("utf-8-sig")


def write_csv(book: "OrderedDict[str, list[list]]", stem: str) -> tuple[bytes, str]:
    """One sheet -> a CSV file; several sheets -> a ZIP of CSVs."""
    if len(book) == 1:
        rows = next(iter(book.values()))
        return _csv_bytes(rows), "csv"
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        used: set = set()
        for name, rows in book.items():
            fname = _safe_sheet_name(name, used)
            zf.writestr(f"{stem} - {fname}.csv", _csv_bytes(rows))
    return buf.getvalue(), "zip"


def write_json(book: "OrderedDict[str, list[list]]") -> bytes:
    """First row of each sheet is used as the keys of each record."""
    out = OrderedDict()
    for name, rows in book.items():
        rows = _trim_rows(rows)
        if not rows:
            out[name] = []
            continue
        header = [_cell_text(h) or f"column_{i + 1}" for i, h in enumerate(rows[0])]
        records = []
        for r in rows[1:]:
            rec = OrderedDict()
            for key, val in zip(header, r):
                val = _clean(val)
                rec[key] = val if isinstance(val, (int, float, bool)) else _cell_text(val)
            records.append(rec)
        out[name] = records
    payload = next(iter(out.values())) if len(out) == 1 else out
    return json.dumps(payload, indent=2, ensure_ascii=False).encode("utf-8")


# --------------------------------------------------------------------------
# Writers: document
# --------------------------------------------------------------------------


def write_docx(blocks: list[tuple]) -> bytes:
    import docx
    from docx.enum.section import WD_ORIENT
    from docx.shared import Pt

    document = docx.Document()
    widest = max((len(b[1][0]) for b in blocks if b[0] == "table" and b[1]), default=0)
    if widest > 6:
        section = document.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    for block in blocks:
        kind = block[0]
        if kind == "heading":
            document.add_heading(block[1], level=min(block[2], 9))
        elif kind == "para":
            document.add_paragraph(block[1])
        elif kind == "table":
            rows = block[1]
            width = max(len(r) for r in rows)
            table = document.add_table(rows=len(rows), cols=width)
            table.style = "Table Grid"
            for i, row in enumerate(rows):
                for j in range(width):
                    cell = table.cell(i, j)
                    cell.text = _cell_text(row[j]) if j < len(row) else ""
                    if i == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
                    if width > 8:
                        for run in cell.paragraphs[0].runs:
                            run.font.size = Pt(8)
            document.add_paragraph()
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def write_pdf(blocks: list[tuple], title: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    widest = max((max(len(r) for r in b[1]) for b in blocks if b[0] == "table" and b[1]), default=0)
    pagesize = landscape(A4) if widest > 6 else A4
    buf = io.BytesIO()
    margin = 15 * mm
    doc = SimpleDocTemplate(
        buf, pagesize=pagesize, leftMargin=margin, rightMargin=margin,
        topMargin=margin, bottomMargin=margin, title=title,
    )
    styles = getSampleStyleSheet()
    usable = pagesize[0] - 2 * margin
    story = []
    for block in blocks:
        kind = block[0]
        if kind == "heading":
            style = styles[f"Heading{min(max(block[2], 1), 6)}"]
            story.append(Paragraph(html.escape(block[1]), style))
        elif kind == "para":
            story.append(Paragraph(html.escape(block[1]).replace("\n", "<br/>"), styles["BodyText"]))
        elif kind == "table":
            rows = block[1]
            width = max(len(r) for r in rows)
            font_size = 9 if width <= 6 else 7 if width <= 12 else 5.5
            cell_style = ParagraphStyle(
                "cell", parent=styles["BodyText"], fontSize=font_size, leading=font_size * 1.2
            )
            data = [
                [Paragraph(html.escape(_cell_text(r[j]) if j < len(r) else ""), cell_style) for j in range(width)]
                for r in rows
            ]
            # Column widths proportional to content length, within the page width.
            lengths = [
                max(4, min(40, max(len(_cell_text(r[j])) if j < len(r) else 0 for r in rows)))
                for j in range(width)
            ]
            total = sum(lengths)
            col_widths = [usable * l / total for l in lengths]
            table = LongTable(data, colWidths=col_widths, repeatRows=1)
            table.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9aa5b1")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8edf3")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ]))
            story.append(table)
            story.append(Spacer(1, 6 * mm))
    if not story:
        story.append(Paragraph("(empty)", styles["BodyText"]))
    doc.build(story)
    return buf.getvalue()


def write_html(blocks: list[tuple], title: str) -> bytes:
    parts = [
        "<!doctype html>",
        '<html><head><meta charset="utf-8">',
        f"<title>{html.escape(title)}</title>",
        "<style>body{font-family:system-ui,sans-serif;max-width:60rem;margin:2rem auto;padding:0 1rem;"
        "line-height:1.5}table{border-collapse:collapse;margin:1rem 0}td,th{border:1px solid #ccc;"
        "padding:.3rem .6rem;text-align:left;vertical-align:top}th{background:#f2f4f7}</style>",
        "</head><body>",
    ]
    for block in blocks:
        if block[0] == "heading":
            lvl = min(max(block[2], 1), 6)
            parts.append(f"<h{lvl}>{html.escape(block[1])}</h{lvl}>")
        elif block[0] == "para":
            parts.append(f"<p>{html.escape(block[1])}</p>")
        elif block[0] == "table":
            rows = block[1]
            parts.append("<table>")
            for i, row in enumerate(rows):
                tag = "th" if i == 0 else "td"
                cells = "".join(f"<{tag}>{html.escape(_cell_text(c))}</{tag}>" for c in row)
                parts.append(f"<tr>{cells}</tr>")
            parts.append("</table>")
    parts.append("</body></html>")
    return "\n".join(parts).encode("utf-8")


def write_markdown(blocks: list[tuple]) -> bytes:
    def esc(text: str) -> str:
        return text.replace("|", "\\|").replace("\n", " ")

    out = []
    for block in blocks:
        if block[0] == "heading":
            out.append("#" * min(max(block[2], 1), 6) + " " + block[1])
        elif block[0] == "para":
            out.append(block[1])
        elif block[0] == "table":
            rows = block[1]
            width = max(len(r) for r in rows)
            norm = [[esc(_cell_text(r[j])) if j < len(r) else "" for j in range(width)] for r in rows]
            out.append("| " + " | ".join(norm[0]) + " |")
            out.append("|" + "---|" * width)
            out.extend("| " + " | ".join(r) + " |" for r in norm[1:])
        out.append("")
    return "\n".join(out).encode("utf-8")


def write_text(blocks: list[tuple]) -> bytes:
    out = []
    for block in blocks:
        if block[0] == "heading":
            out.append(block[1].upper() if block[2] <= 1 else block[1])
            out.append("")
        elif block[0] == "para":
            out.append(block[1])
            out.append("")
        elif block[0] == "table":
            rows = [[_cell_text(c) for c in r] for r in block[1]]
            width = max(len(r) for r in rows)
            rows = [r + [""] * (width - len(r)) for r in rows]
            col_w = [min(40, max(len(r[j]) for r in rows)) for j in range(width)]
            for r in rows:
                out.append("  ".join(c[:40].ljust(col_w[j]) for j, c in enumerate(r)).rstrip())
            out.append("")
    return "\n".join(out).encode("utf-8")


# --------------------------------------------------------------------------
# Entry point
# --------------------------------------------------------------------------


def detect_format(filename: str) -> str:
    ext = os.path.splitext(filename or "")[1].lower().lstrip(".")
    if ext == "xlsm":
        ext = "xlsx"
    if ext not in SPREADSHEET_INPUTS | DOCUMENT_INPUTS | PDF_INPUTS:
        raise ConversionError(
            f"Unsupported file type '.{ext}'. Supported: Excel, CSV, PDF, Word, Google Docs and Sheets."
        )
    return ext


def convert(data: bytes, src: str, dst: str, stem: str = "converted",
            fetch=fetch_google_file) -> Result:
    """Convert ``data`` (in format ``src``) to format ``dst``."""
    if dst not in outputs_for(src):
        raise ConversionError(f"Cannot convert {LABELS.get(src, src)} to {LABELS.get(dst, dst)}.")

    # Google shortcut files / links are resolved into the real export first.
    if src in ("gdoc", "gsheet"):
        text = data.decode("utf-8", errors="ignore").strip()
        url = text if text.startswith("http") else google_url_from_shortcut(data)
        src, data = fetch(url)

    # Direct LibreOffice routes give the best fidelity for Office -> PDF / DOC.
    lo = libreoffice_path() is not None
    if lo and src in ("docx", "doc", "xlsx", "xls") and dst in ("pdf", "doc", "docx") and src != dst:
        if not (src in ("xlsx", "xls") and dst != "pdf"):
            try:
                return _result(libreoffice_convert(data, src, dst), stem, dst)
            except ConversionError:
                pass  # fall back to the pure-Python path below
    if src == "doc":
        data, src = libreoffice_convert(data, "doc", "docx"), "docx"
    if src == dst:
        return _result(data, stem, dst)

    # Load into an intermediate shape.
    if src in SPREADSHEET_INPUTS:
        book = read_workbook(data, src)
        blocks = None
    elif src == "docx":
        blocks = read_docx(data)
        book = None
    else:  # pdf
        if dst in ("xlsx", "csv", "json"):
            try:
                blocks = read_pdf(data)
            except ConversionError:
                blocks = []
            tables = [b for b in blocks if b[0] == "table"]
            book = blocks_to_workbook(tables) if tables else pdf_line_rows(data)
            if not book:
                raise ConversionError("No text or tables found in this PDF.")
        else:
            blocks = read_pdf(data)
            book = None

    # Render.
    if dst in ("xlsx", "csv", "json"):
        if book is None:
            book = blocks_to_workbook(blocks)
        if dst == "xlsx":
            return _result(write_xlsx(book), stem, "xlsx")
        if dst == "csv":
            out, ext = write_csv(book, stem)
            return _result(out, stem, ext)
        return _result(write_json(book), stem, "json")

    if blocks is None:
        blocks = workbook_to_blocks(book)
    if dst == "pdf":
        return _result(write_pdf(blocks, stem), stem, "pdf")
    if dst == "docx":
        return _result(write_docx(blocks), stem, "docx")
    if dst == "doc":
        return _result(libreoffice_convert(write_docx(blocks), "docx", "doc"), stem, "doc")
    if dst == "html":
        return _result(write_html(blocks, stem), stem, "html")
    if dst == "md":
        return _result(write_markdown(blocks), stem, "md")
    if dst == "txt":
        return _result(write_text(blocks), stem, "txt")
    raise ConversionError(f"Unsupported output format: {dst}")


def _result(data: bytes, stem: str, ext: str) -> Result:
    return Result(data=data, filename=f"{stem}.{ext}", mimetype=MIME_TYPES[ext])
